"""
Evaluation script for sf-customer-health-dashboard task.

Checks:
1. Excel file Customer_Health.xlsx with 3 sheets
2. Teamly "Customer Health Dashboard" page created
3. Calendar events for critical accounts with score < 15

A subset of checks are CRITICAL (semantic core deliverables). Any critical
failure forces FAIL regardless of the overall pass rate.
"""
import argparse
import json
import os
import sys
from datetime import date

import openpyxl
import psycopg2

DB_CONFIG = {
    "host": os.environ.get("PGHOST", "localhost"),
    "port": 5432,
    "dbname": os.environ.get("PGDATABASE", "cowork_gym"),
    "user": "eigent",
    "password": "camel",
}

PASS_COUNT = 0
FAIL_COUNT = 0
CRITICAL_FAILURES = []


def check(name, condition, detail="", critical=False):
    global PASS_COUNT, FAIL_COUNT
    tag = "CRITICAL " if critical else ""
    if condition:
        PASS_COUNT += 1
        print(f"  [PASS] {tag}{name}")
    else:
        FAIL_COUNT += 1
        detail_str = f": {detail[:200]}" if detail else ""
        print(f"  [FAIL] {tag}{name}{detail_str}")
        if critical:
            CRITICAL_FAILURES.append(name)


def num_close(a, b, abs_tol=1.0, rel_tol=0.05):
    try:
        a_f, b_f = float(a), float(b)
        return abs(a_f - b_f) <= max(abs_tol, abs(b_f) * rel_tol)
    except (TypeError, ValueError):
        return False


def str_match(a, b):
    if a is None or b is None:
        return a is None and b is None
    return str(a).strip().lower() == str(b).strip().lower()


def compute_expected_values():
    conn = psycopg2.connect(**DB_CONFIG)
    cur = conn.cursor()

    cur.execute("""
        WITH order_stats AS (
            SELECT "CUSTOMER_ID",
                COUNT(*) as order_count,
                MAX("ORDER_DATE") as last_order_date
            FROM sf_data."SALES_DW__PUBLIC__ORDERS"
            GROUP BY "CUSTOMER_ID"
        ),
        health AS (
            SELECT c."CUSTOMER_ID", c."CUSTOMER_NAME", c."SEGMENT", c."REGION",
                ROUND(GREATEST(0, 100 - (CURRENT_DATE - COALESCE(os.last_order_date, c."SIGNUP_DATE")))::numeric, 2) as recency_score,
                ROUND(LEAST(100, COALESCE(os.order_count, 0) * 10)::numeric, 2) as frequency_score,
                ROUND(LEAST(100, c."LIFETIME_VALUE" / 50.0)::numeric, 2) as monetary_score,
                ROUND((0.4 * GREATEST(0, 100 - (CURRENT_DATE - COALESCE(os.last_order_date, c."SIGNUP_DATE"))) +
                0.3 * LEAST(100, COALESCE(os.order_count, 0) * 10) +
                0.3 * LEAST(100, c."LIFETIME_VALUE" / 50.0))::numeric, 2) as health_score
            FROM sf_data."SALES_DW__PUBLIC__CUSTOMERS" c
            LEFT JOIN order_stats os ON c."CUSTOMER_ID" = os."CUSTOMER_ID"
        )
        SELECT "CUSTOMER_NAME", "SEGMENT", "REGION",
            recency_score, frequency_score, monetary_score, health_score,
            CASE WHEN health_score >= 80 THEN 'Healthy'
                 WHEN health_score >= 50 THEN 'At Risk'
                 ELSE 'Critical' END as status
        FROM health
        ORDER BY health_score ASC
    """)
    all_rows = cur.fetchall()

    # Summary by segment
    segments = {}
    for r in all_rows:
        seg = r[1]
        if seg not in segments:
            segments[seg] = {"total": 0, "healthy": 0, "at_risk": 0, "critical": 0, "scores": []}
        segments[seg]["total"] += 1
        segments[seg]["scores"].append(float(r[6]))
        if r[7] == "Healthy":
            segments[seg]["healthy"] += 1
        elif r[7] == "At Risk":
            segments[seg]["at_risk"] += 1
        else:
            segments[seg]["critical"] += 1

    segment_summary = []
    for seg in sorted(segments.keys()):
        s = segments[seg]
        avg = round(sum(s["scores"]) / len(s["scores"]), 2)
        segment_summary.append({
            "Segment": seg,
            "Total_Customers": s["total"],
            "Healthy_Count": s["healthy"],
            "At_Risk_Count": s["at_risk"],
            "Critical_Count": s["critical"],
            "Avg_Health_Score": avg,
        })

    critical_rows = [r for r in all_rows if r[7] == "Critical"]
    below_15 = [r for r in all_rows if float(r[6]) < 15]

    cur.close()
    conn.close()

    return {
        "total_customers": len(all_rows),
        "critical_count": len(critical_rows),
        "at_risk_count": sum(1 for r in all_rows if r[7] == "At Risk"),
        "healthy_count": sum(1 for r in all_rows if r[7] == "Healthy"),
        "segment_summary": segment_summary,
        "below_15": below_15,
        "top5_lowest": all_rows[:5],
    }


def get_sheet(wb, name):
    for s in wb.sheetnames:
        if str_match(s, name):
            return wb[s]
    return None


def check_excel(agent_workspace, expected):
    print("\n=== Checking Excel Output ===")

    agent_file = os.path.join(agent_workspace, "Customer_Health.xlsx")
    check("Excel file exists", os.path.isfile(agent_file), f"Expected {agent_file}")
    if not os.path.isfile(agent_file):
        return

    try:
        wb = openpyxl.load_workbook(agent_file, data_only=True)
    except Exception as e:
        check("Excel file readable", False, str(e))
        return
    check("Excel file readable", True)

    for sn in ["Health Scores", "Critical Accounts", "Summary by Segment"]:
        found = any(str_match(s, sn) for s in wb.sheetnames)
        check(f"Sheet '{sn}' exists", found, f"Found: {wb.sheetnames}")

    # --- Health Scores ---
    print("\n--- Health Scores ---")
    ws = get_sheet(wb, "Health Scores")
    if ws:
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        check("Health Scores total rows",
              num_close(len(rows), expected["total_customers"], 10),
              f"Expected ~{expected['total_customers']}, got {len(rows)}")

        # CRITICAL: top-3 lowest customers' Health_Score must match within +-0.5
        # (verifies the scoring formula was applied correctly on the core rows).
        if len(rows) >= 3:
            for i, exp_row in enumerate(expected["top5_lowest"][:3]):
                agent_row = rows[i]
                got = agent_row[6] if agent_row else None
                check(f"Row {i+1} health score",
                      got is not None and num_close(got, exp_row[6], 0.5),
                      f"Expected ~{exp_row[6]}, got {got}", critical=True)

        # CRITICAL: rows sorted by Health_Score ascending.
        try:
            scores = [float(r[6]) for r in rows if r and r[6] is not None]
            sorted_ok = all(scores[i] <= scores[i + 1] + 0.01 for i in range(len(scores) - 1))
        except (TypeError, ValueError):
            sorted_ok = False
        check("Health Scores sorted by Health_Score ascending", sorted_ok,
              "Rows are not in ascending health-score order", critical=True)

    # --- Critical Accounts ---
    print("\n--- Critical Accounts ---")
    ws = get_sheet(wb, "Critical Accounts")
    if ws:
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        check("Critical Accounts count",
              num_close(len(rows), expected["critical_count"], 20),
              f"Expected ~{expected['critical_count']}, got {len(rows)}")

    # --- Summary by Segment ---
    print("\n--- Summary by Segment ---")
    ws = get_sheet(wb, "Summary by Segment")
    if ws:
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        # CRITICAL: exactly 4 segment rows.
        check("Summary rows", len(rows) == 4,
              f"Expected 4, got {len(rows)}", critical=True)

        for exp_seg in expected["segment_summary"]:
            seg = exp_seg["Segment"]
            matched = None
            for r in rows:
                if r and str_match(r[0], seg):
                    matched = r
                    break
            if matched:
                check(f"{seg} Total_Customers",
                      num_close(matched[1], exp_seg["Total_Customers"], 5),
                      f"Expected {exp_seg['Total_Customers']}, got {matched[1]}")
                check(f"{seg} Healthy_Count",
                      num_close(matched[2], exp_seg["Healthy_Count"], 5),
                      f"Expected {exp_seg['Healthy_Count']}, got {matched[2]}")
                # CRITICAL: classification thresholds applied correctly.
                check(f"{seg} Critical_Count",
                      num_close(matched[4], exp_seg["Critical_Count"], 2),
                      f"Expected {exp_seg['Critical_Count']}, got {matched[4]}",
                      critical=True)
                check(f"{seg} Avg_Health_Score",
                      num_close(matched[5], exp_seg["Avg_Health_Score"], 1.0),
                      f"Expected {exp_seg['Avg_Health_Score']}, got {matched[5]}",
                      critical=True)
            else:
                check(f"Segment '{seg}' found", False, "Not in output", critical=True)


def check_teamly(expected):
    print("\n=== Checking Teamly Dashboard Page ===")
    conn = psycopg2.connect(**DB_CONFIG)
    cur = conn.cursor()

    # Only user-created pages (seed pages have id <= 3).
    cur.execute("""
        SELECT id, title, COALESCE(body, '') FROM teamly.pages
        WHERE id > 3
        ORDER BY id
    """)
    pages = cur.fetchall()
    cur.close()
    conn.close()

    # Title grep broadened to RU+EN: 'customer health dashboard' EN literal,
    # or any RU/EN keyword combination the agent might legitimately use.
    dash = None
    for pid, title, body in pages:
        tl = (title or "").lower()
        if ("health" in tl and "dashboard" in tl) \
                or ("customer" in tl and "dashboard" in tl) \
                or ("клиент" in tl and ("здоров" in tl or "панель" in tl)) \
                or ("панель" in tl and "здоров" in tl):
            dash = (pid, title, body)
            break

    check("Teamly 'Customer Health Dashboard' page exists (RU or EN)",
          dash is not None,
          f"new pages: {[(p[0], p[1]) for p in pages]}", critical=True)

    if dash is None:
        check("Teamly dashboard has content blocks", False, "no dashboard page")
        check("Teamly dashboard contains total count + 5 lowest-score names",
              False, "no dashboard page", critical=True)
        return

    text = ((dash[1] or "") + " " + (dash[2] or "")).lower()
    check("Teamly dashboard has content blocks", len(text.strip()) > 0,
          "page body empty")

    # CRITICAL: dashboard substance — total customer count present AND all five
    # lowest-score customer names present (names are English data literals).
    total = str(expected["total_customers"])
    total_present = total in text
    names = [str(r[0]).lower() for r in expected["top5_lowest"][:5]]
    names_present = sum(1 for n in names if n and n in text)
    check("Teamly dashboard contains total count + 5 lowest-score names",
          total_present and names_present >= 5,
          f"total({total})={total_present}, names {names_present}/5",
          critical=True)


def check_calendar(expected):
    print("\n=== Checking Calendar Events ===")
    conn = psycopg2.connect(**DB_CONFIG)
    cur = conn.cursor()

    cur.execute("""
        SELECT summary, description, start_datetime
        FROM gcal.events
        WHERE LOWER(summary) LIKE '%follow%up%' OR LOWER(summary) LIKE '%follow-up%'
        ORDER BY start_datetime
    """)
    events = cur.fetchall()

    below_15_count = len(expected["below_15"])
    # CRITICAL: exactly one 'Follow-up Call' event per below-15 customer
    # (live-computed count, not a >= lower bound).
    check(f"Calendar 'Follow-up Call' event count equals below-15 count",
          len(events) == below_15_count,
          f"Expected {below_15_count}, got {len(events)}", critical=True)

    # CRITICAL: first 3 events match the lowest-score customers in order.
    if below_15_count >= 3:
        for i in range(3):
            evt_summary = (events[i][0] or "") if i < len(events) else ""
            exp_name = expected["below_15"][i][0]
            check(f"Event {i+1} mentions customer",
                  exp_name.lower() in evt_summary.lower(),
                  f"Expected '{exp_name}' in '{evt_summary}'", critical=True)

    # CRITICAL: events start on or after 2026-03-09 (Monday after analysis).
    if events:
        first_dt = events[0][2]
        if first_dt:
            first_date = first_dt.date() if hasattr(first_dt, 'date') else first_dt
            check("First event on or after 2026-03-09",
                  str(first_date) >= "2026-03-09",
                  f"First event date: {first_date}", critical=True)

    cur.close()
    conn.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--agent_workspace", required=False, default=".")
    parser.add_argument("--groundtruth_workspace", required=False, default=".")
    parser.add_argument("--launch_time", required=False)
    parser.add_argument("--res_log_file", required=False)
    args = parser.parse_args()

    print("=== Computing Expected Values ===")
    try:
        expected = compute_expected_values()
        print(f"  Total customers: {expected['total_customers']}")
        print(f"  Critical: {expected['critical_count']}")
        print(f"  Customers with score < 15: {len(expected['below_15'])}")
    except Exception as e:
        print(f"  ERROR: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)

    check_excel(args.agent_workspace, expected)
    check_teamly(expected)
    check_calendar(expected)

    total = PASS_COUNT + FAIL_COUNT
    pass_rate = PASS_COUNT / total if total > 0 else 0
    critical_ok = len(CRITICAL_FAILURES) == 0
    success = critical_ok and pass_rate >= 0.8

    print(f"\n=== SUMMARY ===")
    print(f"  Passed: {PASS_COUNT}")
    print(f"  Failed: {FAIL_COUNT}")
    print(f"  Pass Rate: {pass_rate:.1%}")
    if not critical_ok:
        print(f"  CRITICAL FAILURES: {CRITICAL_FAILURES}")
    print(f"  Overall: {'PASS' if success else 'FAIL'}")

    if args.res_log_file:
        result = {
            "passed": PASS_COUNT,
            "failed": FAIL_COUNT,
            "pass_rate": round(pass_rate, 3),
            "critical_failures": CRITICAL_FAILURES,
            "success": success,
        }
        with open(args.res_log_file, "w") as f:
            json.dump(result, f, indent=2)

    if not critical_ok:
        print("  FAIL: one or more CRITICAL checks failed.")
        sys.exit(1)
    sys.exit(0 if success else 1)


if __name__ == "__main__":
    main()
