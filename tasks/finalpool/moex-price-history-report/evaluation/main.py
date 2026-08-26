"""Evaluation for yf-price-history-report."""
import argparse
import os
import sys
import openpyxl


def num_close(a, b, tol=1.0):
    if a is None and b is None:
        return True
    if a is None or b is None:
        return False
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return str(a).strip().lower() == str(b).strip().lower()


def str_match(a, b):
    if a is None or b is None:
        return a is None and b is None
    return str(a).strip().lower() == str(b).strip().lower()


def load_sheet_rows(wb, sheet_name):
    for name in wb.sheetnames:
        if name.strip().lower() == sheet_name.strip().lower():
            return [[cell.value for cell in row] for row in wb[name].iter_rows()]
    return None


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--agent_workspace", required=False)
    parser.add_argument("--groundtruth_workspace", required=False)
    parser.add_argument("--launch_time", required=False)
    parser.add_argument("--res_log_file", required=False)
    args = parser.parse_args()

    task_root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    gt_dir = args.groundtruth_workspace or os.path.join(task_root, "groundtruth_workspace")

    agent_file = os.path.join(args.agent_workspace, "YF_Price_Report.xlsx")
    gt_file = os.path.join(gt_dir, "YF_Price_Report.xlsx")

    if not os.path.exists(agent_file):
        print(f"FAIL: Agent output not found: {agent_file}")
        sys.exit(1)
    if not os.path.exists(gt_file):
        print(f"FAIL: Groundtruth not found: {gt_file}")
        sys.exit(1)

    agent_wb = openpyxl.load_workbook(agent_file, data_only=True)
    gt_wb = openpyxl.load_workbook(gt_file, data_only=True)

    all_errors = []
    critical_failures = []

    # The 6 seeded MOEX tickers the agent MUST track (no US/yahoo-finance leakage).
    EXPECTED_SYMBOLS = {"sber.me", "gazp.me", "lkoh.me", "tcsg.me", "mgnt.me", "mtss.me"}

    # Check sheet: Price Summary
    print(f"  Checking Price Summary...")
    a_rows = load_sheet_rows(agent_wb, "Price Summary")
    g_rows = load_sheet_rows(gt_wb, "Price Summary")
    if a_rows is None:
        all_errors.append("Sheet 'Price Summary' not found in agent output")
    elif g_rows is None:
        all_errors.append("Sheet 'Price Summary' not found in groundtruth")
    else:
        sheet_name = "Price Summary"
        errors = []
        a_data = a_rows[1:] if len(a_rows) > 1 else []
        g_data = g_rows[1:] if len(g_rows) > 1 else []
        
        a_lookup = {}
        for row in a_data:
            if row and row[0] is not None:
                a_lookup[str(row[0]).strip().lower()] = row

        # CRITICAL: exact seeded-symbol universe — all 6 MOEX tickers, no extra US symbols.
        agent_symbols = set(a_lookup.keys())
        if agent_symbols != EXPECTED_SYMBOLS:
            missing = sorted(EXPECTED_SYMBOLS - agent_symbols)
            extra = sorted(agent_symbols - EXPECTED_SYMBOLS)
            critical_failures.append(
                f"Price Summary symbol set mismatch: missing={missing} extra={extra}"
            )

        for g_row in g_data:
            if not g_row or g_row[0] is None:
                continue
            key = str(g_row[0]).strip().lower()
            a_row = a_lookup.get(key)
            if a_row is None:
                errors.append(f"Missing row: {g_row[0]}")
                continue

            if len(a_row) > 1 and len(g_row) > 1:
                if not num_close(a_row[1], g_row[1], 5.0):
                    errors.append(f"{key}.Avg_Close: {a_row[1]} vs {g_row[1]} (tol=5.0)")
                    critical_failures.append(f"{key}.Avg_Close: {a_row[1]} vs {g_row[1]} (tol=5.0)")

            if len(a_row) > 2 and len(g_row) > 2:
                if not num_close(a_row[2], g_row[2], 5.0):
                    errors.append(f"{key}.Min_Close: {a_row[2]} vs {g_row[2]} (tol=5.0)")
                    critical_failures.append(f"{key}.Min_Close: {a_row[2]} vs {g_row[2]} (tol=5.0)")

            if len(a_row) > 3 and len(g_row) > 3:
                if not num_close(a_row[3], g_row[3], 5.0):
                    errors.append(f"{key}.Max_Close: {a_row[3]} vs {g_row[3]} (tol=5.0)")
                    critical_failures.append(f"{key}.Max_Close: {a_row[3]} vs {g_row[3]} (tol=5.0)")

            if len(a_row) > 4 and len(g_row) > 4:
                if not num_close(a_row[4], g_row[4], 5):
                    errors.append(f"{key}.Trading_Days: {a_row[4]} vs {g_row[4]} (tol=5)")

            if len(a_row) > 5 and len(g_row) > 5:
                if not num_close(a_row[5], g_row[5], 2.0):
                    errors.append(f"{key}.Volatility: {a_row[5]} vs {g_row[5]} (tol=2.0)")
                    critical_failures.append(f"{key}.Volatility: {a_row[5]} vs {g_row[5]} (tol=2.0)")
        if errors:
            all_errors.extend(errors)
            print(f"    ERRORS: {len(errors)}")
            for e in errors[:5]:
                print(f"      {e}")
        else:
            print(f"    PASS")


    # Check sheet: Summary
    print(f"  Checking Summary...")
    a_rows = load_sheet_rows(agent_wb, "Summary")
    g_rows = load_sheet_rows(gt_wb, "Summary")
    if a_rows is None:
        all_errors.append("Sheet 'Summary' not found in agent output")
    elif g_rows is None:
        all_errors.append("Sheet 'Summary' not found in groundtruth")
    else:
        sheet_name = "Summary"
        errors = []
        a_data = a_rows[1:] if len(a_rows) > 1 else []
        g_data = g_rows[1:] if len(g_rows) > 1 else []
        
        a_lookup = {}
        for row in a_data:
            if row and row[0] is not None:
                a_lookup[str(row[0]).strip().lower()] = row
        for g_row in g_data:
            if not g_row or g_row[0] is None:
                continue
            key = str(g_row[0]).strip().lower()
            a_row = a_lookup.get(key)
            if a_row is None:
                errors.append(f"Missing row: {g_row[0]}")
                continue
            
            if len(a_row) > 1 and len(g_row) > 1:
                if not str_match(a_row[1], g_row[1]):
                    errors.append(f"{key}.Value: '{a_row[1]}' vs '{g_row[1]}'")
                    # CRITICAL: Stocks_Tracked, Most_Volatile, Least_Volatile are the
                    # core semantic signals — exact match against the moex seed.
                    if key in ("stocks_tracked", "most_volatile", "least_volatile"):
                        critical_failures.append(
                            f"Summary {g_row[0]}: '{a_row[1]}' vs '{g_row[1]}'"
                        )
        if errors:
            all_errors.extend(errors)
            print(f"    ERRORS: {len(errors)}")
            for e in errors[:5]:
                print(f"      {e}")
        else:
            print(f"    PASS")

    
    docx_path = os.path.join(args.agent_workspace, "Market_Analysis.docx")
    if not os.path.exists(docx_path):
        all_errors.append("Market_Analysis.docx not found")
    else:
        try:
            from docx import Document as _DocCheck
            _doc = _DocCheck(docx_path)
            _text = " ".join(p.text for p in _doc.paragraphs).lower()
            _headings = " ".join(p.text for p in _doc.paragraphs if p.style.name.startswith("Heading")).lower()
            if len(_text.strip()) < 50:
                all_errors.append("Market_Analysis.docx has too little text content (< 50 chars)")
            # Accept RU or EN vocabulary (agent legitimately writes Russian narrative).
            _kws = ["market", "analysis", "рынок", "рынк", "анализ", "волатильн"]
            _missing = [k for k in _kws if k not in _text and k not in _headings]
            if len(_missing) == len(_kws):
                all_errors.append(f"Market_Analysis.docx missing expected keywords: {_missing}")
        except ImportError:
            if os.path.getsize(docx_path) < 100:
                all_errors.append("Market_Analysis.docx too small")
        except Exception as _e:
            all_errors.append(f"Error reading Market_Analysis.docx: {_e}")

    # CRITICAL GATE: any critical semantic failure => hard fail before final decision.
    if critical_failures:
        print(f"\n=== CRITICAL FAILURE ({len(critical_failures)}) ===")
        for e in critical_failures[:10]:
            print(f"  CRITICAL: {e}")
        sys.exit(1)

    if all_errors:
        print(f"\n=== RESULT: FAIL ({len(all_errors)} errors) ===")
        for e in all_errors[:10]:
            print(f"  {e}")
        sys.exit(1)
    else:
        print("\n=== RESULT: PASS ===")
        sys.exit(0)


if __name__ == "__main__":
    main()
