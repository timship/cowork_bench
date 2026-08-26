"""
Evaluation for canvas-grade-distribution-word.
Checks:
1. Excel file Grade_Data.xlsx with correct grade distribution and assignment averages
2. Word document Grade_Report.docx with required content
"""
import argparse
import os
import sys

import openpyxl
import psycopg2

DB_CONFIG = {
    "host": os.environ.get("PGHOST", "localhost"),
    "port": 5432,
    "dbname": "cowork_gym",
    "user": "eigent",
    "password": "camel",
}

COURSE_ID = 16  # FFF-2013J

PASS_COUNT = 0
FAIL_COUNT = 0
CRITICAL_FAILS = []

# Names of CHECKS that must pass; any critical fail => FAIL regardless of accuracy.
CRITICAL_CHECKS = {
    "Distribution band A Count",
    "Distribution band B Count",
    "Distribution band C Count",
    "Distribution band D Count",
    "Distribution band F Count",
    "Distribution band A Percentage",
    "Distribution band B Percentage",
    "Distribution band C Percentage",
    "Distribution band D Percentage",
    "Distribution band F Percentage",
    "All assignment averages correct",
    "Word doc mentions FFF-2013J",
    "Word doc mentions total submissions",
}


def record(name, passed, detail="", critical=False):
    global PASS_COUNT, FAIL_COUNT
    crit = critical or (name in CRITICAL_CHECKS)
    tag = "CRITICAL " if crit else ""
    if passed:
        PASS_COUNT += 1
        print(f"  [PASS] {tag}{name}")
    else:
        FAIL_COUNT += 1
        msg = f": {detail[:300]}" if detail else ""
        print(f"  [FAIL] {tag}{name}{msg}")
        if crit:
            CRITICAL_FAILS.append(name)


def num_close(a, b, tol=1.0):
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return False


def str_match(a, b):
    if a is None or b is None:
        return a is None and b is None
    return str(a).strip().lower() == str(b).strip().lower()


def load_sheet_rows(wb, sheet_name):
    for name in wb.sheetnames:
        if name.strip().lower() == sheet_name.strip().lower():
            return [[cell.value for cell in row] for row in wb[name].iter_rows()]
    return None


def get_expected_data():
    """Compute expected grade distribution from Canvas DB."""
    conn = psycopg2.connect(**DB_CONFIG)
    cur = conn.cursor()

    cur.execute("""
        SELECT s.score, a.points_possible
        FROM canvas.submissions s
        JOIN canvas.assignments a ON s.assignment_id = a.id
        WHERE a.course_id = %s AND s.score IS NOT NULL
    """, (COURSE_ID,))
    submissions = cur.fetchall()

    bands = {'A': 0, 'B': 0, 'C': 0, 'D': 0, 'F': 0}
    for score, points in submissions:
        if points and float(points) > 0:
            pct = (float(score) / float(points)) * 100
        else:
            pct = float(score)
        if pct >= 90:
            bands['A'] += 1
        elif pct >= 80:
            bands['B'] += 1
        elif pct >= 70:
            bands['C'] += 1
        elif pct >= 60:
            bands['D'] += 1
        else:
            bands['F'] += 1

    total = sum(bands.values())

    cur.execute("""
        SELECT a.name,
               ROUND(AVG(s.score)::numeric, 2) as avg_score,
               COUNT(s.id) as submission_count
        FROM canvas.assignments a
        JOIN canvas.submissions s ON s.assignment_id = a.id
        WHERE a.course_id = %s AND s.score IS NOT NULL
        GROUP BY a.id, a.name
        ORDER BY a.name
    """, (COURSE_ID,))
    assignments = cur.fetchall()

    cur.close()
    conn.close()

    return bands, total, assignments


def check_excel(agent_workspace, gt_workspace, expected_bands, expected_total, expected_assignments):
    """Check Grade_Data.xlsx."""
    print("\n=== Checking Excel ===")
    agent_file = os.path.join(agent_workspace, "Grade_Data.xlsx")
    gt_file = os.path.join(gt_workspace, "Grade_Data.xlsx")

    if not os.path.exists(agent_file):
        record("Excel file exists", False, f"Not found: {agent_file}")
        return False
    record("Excel file exists", True)

    agent_wb = openpyxl.load_workbook(agent_file, data_only=True)
    # gt_file kept for shape reference; expected values are recomputed live from the DB.
    _ = gt_file

    # Check Distribution sheet
    a_rows = load_sheet_rows(agent_wb, "Distribution")

    if a_rows is None:
        record("Sheet 'Distribution' exists", False, f"Sheets: {agent_wb.sheetnames}")
        return False
    record("Sheet 'Distribution' exists", True)

    a_data = a_rows[1:] if len(a_rows) > 1 else []

    record("Distribution has 5 rows", len(a_data) == 5, f"Got {len(a_data)}")

    a_lookup = {}
    for row in a_data:
        if row and row[0]:
            a_lookup[str(row[0]).strip().upper()] = row

    # Use live-DB-computed expected bands/total (deterministic) as the source of truth.
    for grade in ['A', 'B', 'C', 'D', 'F']:
        exp_count = expected_bands.get(grade, 0)
        exp_pct = round((exp_count / expected_total * 100), 2) if expected_total else 0.0
        a_row = a_lookup.get(grade)
        if a_row is None:
            record(f"Grade band '{grade}' found", False, "Missing")
            record(f"Distribution band {grade} Count", False, "Row missing")
            record(f"Distribution band {grade} Percentage", False, "Row missing")
            continue
        record(f"Grade band '{grade}' found", True)
        if len(a_row) > 1:
            record(f"Distribution band {grade} Count",
                   num_close(a_row[1], exp_count, 1),
                   f"Agent={a_row[1]}, expected={exp_count}")
        else:
            record(f"Distribution band {grade} Count", False, "No Count column")
        if len(a_row) > 2:
            record(f"Distribution band {grade} Percentage",
                   num_close(a_row[2], exp_pct, 0.5),
                   f"Agent={a_row[2]}, expected={exp_pct}")
        else:
            record(f"Distribution band {grade} Percentage", False, "No Percentage column")

    # Check Assignment Averages sheet
    a_rows2 = load_sheet_rows(agent_wb, "Assignment Averages")
    if a_rows2 is None:
        record("Sheet 'Assignment Averages' exists", False, f"Sheets: {agent_wb.sheetnames}")
        return False
    record("Sheet 'Assignment Averages' exists", True)

    a_data2 = a_rows2[1:] if len(a_rows2) > 1 else []

    # expected_assignments is computed live from the Canvas DB: (name, avg_score, count)
    exp_names = [str(a[0]).strip() for a in expected_assignments]
    record("Assignment Averages row count",
           abs(len(a_data2) - len(exp_names)) <= 2,
           f"Expected {len(exp_names)}, got {len(a_data2)}")

    a_assign_lookup = {}
    for row in a_data2:
        if row and row[0]:
            a_assign_lookup[str(row[0]).strip().lower()] = row

    # Check EVERY assignment (not just first 5) against live-DB expected Avg_Score.
    all_assignments_ok = True
    for exp in expected_assignments:
        name = str(exp[0]).strip()
        exp_avg = exp[1]
        a_row = a_assign_lookup.get(name.lower())
        if a_row is None:
            record(f"Assignment '{name}' found", False, "Missing")
            all_assignments_ok = False
            continue
        record(f"Assignment '{name}' found", True)
        ok = len(a_row) > 1 and num_close(a_row[1], exp_avg, 2.0)
        record(f"  Avg_Score '{name}'", ok,
               f"Agent={a_row[1] if len(a_row) > 1 else None}, expected={exp_avg}")
        if not ok:
            all_assignments_ok = False

    # Aggregate CRITICAL check: all assignment averages present and correct, sorted alpha.
    agent_names = [str(r[0]).strip() for r in a_data2 if r and r[0]]
    sorted_ok = agent_names == sorted(agent_names, key=lambda s: s.lower())
    record("All assignment averages correct",
           all_assignments_ok and sorted_ok,
           f"averages_ok={all_assignments_ok}, sorted={sorted_ok}")

    return True


def check_word(agent_workspace, expected_total, expected_bands):
    """Check Grade_Report.docx exists and has required content."""
    print("\n=== Checking Word Document ===")

    docx_path = os.path.join(agent_workspace, "Grade_Report.docx")
    if not os.path.exists(docx_path):
        record("Word file exists", False, f"Not found: {docx_path}")
        return False
    record("Word file exists", True)

    try:
        from docx import Document
        doc = Document(docx_path)

        all_text = " ".join(p.text for p in doc.paragraphs).lower()
        tables_text = " ".join(
            " ".join(cell.text for cell in row.cells)
            for t in doc.tables for row in t.rows
        ).lower()
        full_text = all_text + " " + tables_text

        # Course code stays English per task.md heading requirement.
        record("Word doc mentions FFF-2013J",
               "fff-2013j" in full_text or "fff 2013j" in full_text,
               "No reference to course code", critical=True)

        # The heading literal is required English, but the agent legitimately writes
        # Russian prose ('Распределение оценок'), so accept RU OR EN wording.
        en_ok = "grade" in full_text and "distribution" in full_text
        ru_ok = "оцен" in full_text and "распределен" in full_text
        record("Document mentions grade distribution",
               en_ok or ru_ok,
               "Missing grade-distribution wording (EN or RU)")

        record("Document has at least one table",
               len(doc.tables) >= 1,
               f"Found {len(doc.tables)} tables")

        # Total graded submissions (live-computed) must appear in prose or table.
        record("Word doc mentions total submissions",
               str(expected_total) in full_text,
               f"Expected {expected_total} to be mentioned", critical=True)

        # The grade-distribution table band rows must match the Distribution counts.
        bands_in_table = sum(
            1 for c in expected_bands.values() if str(c) in tables_text
        )
        record("Word table band counts match distribution",
               bands_in_table >= 4,
               f"Only {bands_in_table}/5 band counts found in tables")

    except ImportError:
        record("python-docx available", False, "Cannot import docx module")
        return False
    except Exception as e:
        record("Word document readable", False, str(e))
        return False

    return True


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--agent_workspace", required=False, default=".")
    parser.add_argument("--groundtruth_workspace", required=False)
    parser.add_argument("--launch_time", required=False)
    parser.add_argument("--res_log_file", required=False)
    args = parser.parse_args()

    task_root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    gt_dir = args.groundtruth_workspace or os.path.join(task_root, "groundtruth_workspace")

    bands, total, assignments = get_expected_data()
    print(f"[eval] Total submissions: {total}")
    print(f"[eval] Grade bands: {bands}")

    check_excel(args.agent_workspace, gt_dir, bands, total, assignments)
    check_word(args.agent_workspace, total, bands)

    total_checks = PASS_COUNT + FAIL_COUNT
    accuracy = (PASS_COUNT / total_checks * 100) if total_checks else 0.0

    print(f"\n=== SUMMARY ===")
    print(f"  Passed: {PASS_COUNT}")
    print(f"  Failed: {FAIL_COUNT}")
    print(f"  Accuracy: {accuracy:.1f}%")

    # CRITICAL gate: any critical failure => hard FAIL regardless of accuracy.
    if CRITICAL_FAILS:
        print(f"  CRITICAL FAILURES: {CRITICAL_FAILS}")
        print(f"  Overall: FAIL (critical check failed)")
        sys.exit(1)

    overall = PASS_COUNT > 0 and accuracy >= 70
    print(f"  Overall: {'PASS' if overall else 'FAIL'}")
    sys.exit(0 if overall else 1)


if __name__ == "__main__":
    main()
