import argparse
import json
import os
import sys
import psycopg2

DB_CONFIG = {
    "host": os.environ.get("PGHOST", "localhost"),
    "port": 5432,
    "dbname": os.environ.get("PGDATABASE", "cowork_gym"),
    "user": "eigent",
    "password": "camel",
}

PASS_COUNT = 0
FAIL_COUNT = 0
FAILED_NAMES = []

# Critical SEMANTIC checks: any failure => overall FAIL regardless of accuracy.
# These verify the substance of the deliverable (correct scores, thresholds,
# routing) rather than mere structural presence.
CRITICAL_CHECKS = {
    "Alignment scores follow rubric (0-100, rounded 1dp)",
    "Status thresholding correct in GSheet (>=30 Aligned, <30 Review Needed)",
    "summary_stats numeric correctness (total/avg/max/min/below30)",
    "Department-head email routing correct (no cross-department leakage)",
    "Dean email body has correct faculty count, average and review count",
}


def check(name, condition, detail=""):
    global PASS_COUNT, FAIL_COUNT
    if condition:
        PASS_COUNT += 1
        print(f"  [PASS] {name}")
    else:
        FAIL_COUNT += 1
        FAILED_NAMES.append(name)
        print(f"  [FAIL] {name}: {str(detail)[:200]}")


def num_close(a, b, tol=5.0):
    try:
        return abs(float(a) - float(b)) <= tol
    except:
        return False


_FALLBACK_FACULTY = ["Dr. Sarah Chen", "Dr. James Okafor", "Dr. Maria Gonzalez", "Dr. Raj Patel"]
_FALLBACK_DEPARTMENTS = ["Biochemistry", "Bioinformatics"]


def _get_faculty_from_roster():
    """Read faculty_roster.csv to get faculty names and departments dynamically."""
    try:
        import csv as _csv
        roster_path = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "initial_workspace",
            "faculty_roster.csv",
        )
        faculty = []
        departments = set()
        with open(roster_path) as f:
            reader = _csv.DictReader(f)
            for row in reader:
                name = row.get("faculty_name", row.get("name", "")).strip()
                dept = row.get("department", "").strip()
                if name:
                    faculty.append(name)
                if dept:
                    departments.add(dept)
        if faculty:
            return faculty, sorted(departments)
        return _FALLBACK_FACULTY, _FALLBACK_DEPARTMENTS
    except Exception:
        return _FALLBACK_FACULTY, _FALLBACK_DEPARTMENTS


def check_xlsx_content(workspace):
    """Check Alignment_Summary.xlsx has valid content."""
    print("\n=== Checking XLSX Content ===")
    try:
        import openpyxl
    except ImportError:
        check("openpyxl available", False, "Cannot import openpyxl")
        return False

    xlsx_path = os.path.join(workspace, "Alignment_Summary.xlsx")
    if not os.path.isfile(xlsx_path):
        check("Alignment_Summary.xlsx exists", False, f"Not found: {xlsx_path}")
        return False
    check("Alignment_Summary.xlsx exists", True)

    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        check("XLSX has at least one sheet", len(wb.worksheets) >= 1,
              f"Found {len(wb.worksheets)} sheets")
        all_ok = True
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            has_data = len(rows) >= 2
            check(f"XLSX sheet '{ws.title}' has data rows", has_data,
                  f"Only {len(rows)} rows")
            if not has_data:
                all_ok = False
        wb.close()
        return all_ok
    except Exception as e:
        check("XLSX readable", False, str(e))
        return False


FACULTY, DEPARTMENTS = _get_faculty_from_roster()
FACULTY_COUNT = len(FACULTY)

# Noise faculty/departments that should NOT appear in outputs.
# NOTE: kept English. The GSheet Department column is populated from the
# (English) faculty roster, not from the russified canvas data, so these
# English guard strings remain valid. They are deliberately NOT russified:
# "Информатика" (Computer Science) is a substring of the legitimate russified
# course name "Биоинформатика", which would cause a false positive.
NOISE_DEPARTMENTS = ["Computer Science", "Mathematics", "Physics", "History"]


def check_reverse_validation():
    print("\n=== Reverse Validation ===")
    conn = psycopg2.connect(**DB_CONFIG)
    cur = conn.cursor()
    try:
        # Check GSheet does not contain irrelevant/noise faculty
        cur.execute("""
            SELECT id FROM gsheet.spreadsheets
            WHERE title ILIKE '%%Research_Teaching_Alignment%%'
               OR title ILIKE '%%Research%%Teaching%%Alignment%%'
            LIMIT 1
        """)
        row = cur.fetchone()
        if row:
            ss_id = row[0]
            cur.execute("""
                SELECT value FROM gsheet.cells
                WHERE spreadsheet_id = %s AND row_index > 0
            """, (ss_id,))
            all_values = " ".join([r[0] for r in cur.fetchall() if r[0]]).lower()
            # Only 4 faculty should be present; check no noise departments
            for dept in NOISE_DEPARTMENTS:
                check(f"GSheet does not contain noise department '{dept}'",
                      dept.lower() not in all_values,
                      f"Found '{dept}' in GSheet data")
            # Check no more than 4 distinct faculty rows
            cur.execute("""
                SELECT COUNT(DISTINCT row_index) FROM gsheet.cells
                WHERE spreadsheet_id = %s AND row_index > 0
            """, (ss_id,))
            data_rows = cur.fetchone()[0]
            check("GSheet has no extra noise rows beyond 4 faculty",
                  data_rows <= 4, f"Found {data_rows} data rows")

        # Check no emails sent to wrong recipients
        noise_emails = ["newsletter@university.edu", "all-faculty@university.edu",
                        "registrar@university.edu"]
        for addr in noise_emails:
            cur.execute(
                "SELECT COUNT(*) FROM email.messages WHERE to_addr::text ILIKE %s",
                (f"%{addr}%",),
            )
            cnt = cur.fetchone()[0]
            check(f"No email sent to noise recipient {addr}", cnt == 0,
                  f"Found {cnt} emails to {addr}")
    except Exception as e:
        check("Reverse validation", False, str(e))
    finally:
        cur.close()
        conn.close()


def _extract_score(entry):
    """Pull a numeric alignment score out of a faculty entry (dict or number)."""
    if isinstance(entry, (int, float)):
        return float(entry)
    if isinstance(entry, dict):
        for k in ("alignment_score", "score", "alignment_score_percent",
                  "alignment", "alignment_percent"):
            if k in entry:
                try:
                    return float(entry[k])
                except (TypeError, ValueError):
                    pass
        # fall back: first numeric value
        for v in entry.values():
            if isinstance(v, (int, float)):
                return float(v)
    return None


def _scores_by_faculty(as_data):
    """Return {faculty_name: score} from alignment_scores.json (list or dict)."""
    out = {}
    if isinstance(as_data, dict):
        for k, v in as_data.items():
            s = _extract_score(v)
            if s is not None:
                out[k] = s
    elif isinstance(as_data, list):
        for e in as_data:
            if isinstance(e, dict):
                name = e.get("faculty_name", e.get("name", ""))
                s = _extract_score(e)
                if name and s is not None:
                    out[name] = s
    return out


def check_score_semantics(as_data):
    """CRITICAL: scores are valid percentages per the rubric (0..100, 1dp)."""
    print("\n=== Checking alignment-score semantics (CRITICAL) ===")
    scores = _scores_by_faculty(as_data)
    have_all = len(scores) == FACULTY_COUNT and all(
        any(fac.lower() in name.lower() for name in scores) for fac in FACULTY
    )
    if not have_all:
        check("Alignment scores follow rubric (0-100, rounded 1dp)", False,
              f"Could not map {FACULTY_COUNT} faculty scores: {scores}")
        return scores
    ok = True
    for name, s in scores.items():
        if not (0.0 <= s <= 100.0):
            ok = False
        # rubric: rounded to one decimal place
        if abs(round(s, 1) - s) > 1e-6:
            ok = False
    check("Alignment scores follow rubric (0-100, rounded 1dp)", ok,
          f"Scores: {scores}")
    return scores


def check_status_thresholding(cur, scores):
    """CRITICAL: every GSheet row Status matches the 30% rubric threshold."""
    print("\n=== Checking GSheet Status thresholding (CRITICAL) ===")
    cur.execute("""
        SELECT id FROM gsheet.spreadsheets
        WHERE title ILIKE '%%Research_Teaching_Alignment%%'
           OR title ILIKE '%%Research%%Teaching%%Alignment%%'
        LIMIT 1
    """)
    row = cur.fetchone()
    if not row:
        check("Status thresholding correct in GSheet (>=30 Aligned, <30 Review Needed)",
              False, "No spreadsheet found")
        return
    ss_id = row[0]
    cur.execute("""
        SELECT id FROM gsheet.sheets
        WHERE spreadsheet_id = %s AND title ILIKE '%%Alignment%%Matrix%%' LIMIT 1
    """, (ss_id,))
    sr = cur.fetchone()
    if not sr:
        check("Status thresholding correct in GSheet (>=30 Aligned, <30 Review Needed)",
              False, "No Alignment Matrix sheet")
        return
    sheet_id = sr[0]
    # Build per-row (row_index -> {col_index: value})
    cur.execute("""
        SELECT row_index, col_index, value FROM gsheet.cells
        WHERE spreadsheet_id = %s AND sheet_id = %s
        ORDER BY row_index, col_index
    """, (ss_id, sheet_id))
    grid = {}
    for ri, ci, val in cur.fetchall():
        grid.setdefault(ri, {})[ci] = val or ""
    if 0 not in grid:
        check("Status thresholding correct in GSheet (>=30 Aligned, <30 Review Needed)",
              False, "No header row")
        return
    header = {ci: (v or "").lower() for ci, v in grid[0].items()}
    score_col = next((ci for ci, h in header.items()
                      if "alignment" in h or "score" in h or "балл" in h), None)
    status_col = next((ci for ci, h in header.items()
                       if "status" in h or "статус" in h), None)
    if score_col is None or status_col is None:
        check("Status thresholding correct in GSheet (>=30 Aligned, <30 Review Needed)",
              False, f"Missing score/status column; headers={header}")
        return
    ok = True
    detail = ""
    checked = 0
    for ri, cols in grid.items():
        if ri == 0:
            continue
        raw = str(cols.get(score_col, "")).replace("%", "").strip().replace(",", ".")
        try:
            sc = float(raw)
        except ValueError:
            ok = False
            detail = f"row {ri} score unparseable: {cols.get(score_col)!r}"
            continue
        status = str(cols.get(status_col, "")).lower()
        is_aligned = ("aligned" in status or "согласов" in status) and "review" not in status and "пересмотр" not in status
        is_review = "review" in status or "пересмотр" in status
        checked += 1
        if sc >= 30:
            if not is_aligned:
                ok = False
                detail = f"row {ri}: score {sc} should be Aligned, got {status!r}"
        else:
            if not is_review:
                ok = False
                detail = f"row {ri}: score {sc} should be Review Needed, got {status!r}"
    if checked == 0:
        ok = False
        detail = "no data rows to validate"
    check("Status thresholding correct in GSheet (>=30 Aligned, <30 Review Needed)",
          ok, detail)


def check_summary_semantics(ss_data, scores):
    """CRITICAL: summary_stats numbers recomputed from alignment scores."""
    print("\n=== Checking summary_stats numeric correctness (CRITICAL) ===")
    vals = list(scores.values())
    if len(vals) != FACULTY_COUNT:
        check("summary_stats numeric correctness (total/avg/max/min/below30)",
              False, f"Need {FACULTY_COUNT} scores, got {len(vals)}")
        return
    exp_total = FACULTY_COUNT
    exp_avg = sum(vals) / len(vals)
    exp_max = max(vals)
    exp_min = min(vals)
    exp_below = sum(1 for v in vals if v < 30)

    def _g(*keys):
        for k in list(ss_data.keys()):
            kn = k.replace("_", "").replace(" ", "").lower()
            for want in keys:
                if want.replace("_", "").replace(" ", "").lower() == kn:
                    return ss_data[k]
        return None

    g_total = _g("total_faculty", "total", "faculty_count", "count")
    g_avg = _g("average_alignment_score", "average_score", "avg", "mean", "average")
    g_max = _g("max_score", "max", "maximum")
    g_min = _g("min_score", "min", "minimum")
    g_below = _g("below_30", "below30", "count_below_30", "review_count",
                 "flagged", "needs_review", "below_threshold")

    ok = True
    detail = []
    if not num_close(g_total, exp_total, tol=0):
        ok = False; detail.append(f"total {g_total}!={exp_total}")
    if not num_close(g_avg, exp_avg, tol=1.0):
        ok = False; detail.append(f"avg {g_avg}!={exp_avg:.2f}")
    if not num_close(g_max, exp_max, tol=0.5):
        ok = False; detail.append(f"max {g_max}!={exp_max}")
    if not num_close(g_min, exp_min, tol=0.5):
        ok = False; detail.append(f"min {g_min}!={exp_min}")
    if g_below is not None and not num_close(g_below, exp_below, tol=0):
        ok = False; detail.append(f"below30 {g_below}!={exp_below}")
    check("summary_stats numeric correctness (total/avg/max/min/below30)",
          ok, "; ".join(detail))


def check_email_routing_semantics(cur, scores):
    """CRITICAL: dept-head emails go to the right head with only that dept's faculty."""
    print("\n=== Checking department-head email routing (CRITICAL) ===")
    # Map dept -> (head_email, [faculty names]) from the roster.
    import csv as _csv
    roster_path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "initial_workspace", "faculty_roster.csv",
    )
    dept_map = {}
    try:
        with open(roster_path) as f:
            for r in _csv.DictReader(f):
                dept = (r.get("department") or "").strip()
                name = (r.get("faculty_name") or r.get("name") or "").strip()
                head = (r.get("dept_head_email") or "").strip().lower()
                if dept:
                    dept_map.setdefault(dept, {"head": head, "faculty": []})
                    if name:
                        dept_map[dept]["faculty"].append(name)
                    if head:
                        dept_map[dept]["head"] = head
    except Exception as e:
        check("Department-head email routing correct (no cross-department leakage)",
              False, f"roster read failed: {e}")
        return

    cur.execute("""
        SELECT to_addr, body_text FROM email.messages
        WHERE subject ILIKE '%%Department%%Research%%Teaching%%Alignment%%Update%%'
    """)
    dept_emails = cur.fetchall()
    ok = True
    detail = []
    for dept, info in dept_map.items():
        head = info["head"]
        own = info["faculty"]
        # find the email addressed to this head
        match = None
        for to_addr, body in dept_emails:
            if head and head in str(to_addr).lower():
                match = (to_addr, body or "")
                break
        if not match:
            ok = False; detail.append(f"no email to {head} for {dept}")
            continue
        body_l = match[1].lower()
        # own faculty must be present
        for fac in own:
            last = fac.split()[-1].lower()
            if last not in body_l:
                ok = False; detail.append(f"{dept} email missing {fac}")
        # other-department faculty must NOT leak in
        for odept, oinfo in dept_map.items():
            if odept == dept:
                continue
            for ofac in oinfo["faculty"]:
                olast = ofac.split()[-1].lower()
                if olast in body_l and olast not in [f.split()[-1].lower() for f in own]:
                    ok = False; detail.append(f"{dept} email leaks {ofac}")
    check("Department-head email routing correct (no cross-department leakage)",
          ok, "; ".join(detail))


def check_dean_email_semantics(cur, scores):
    """CRITICAL: dean email body carries the correct count/average/review count."""
    print("\n=== Checking dean email content (CRITICAL) ===")
    vals = list(scores.values())
    cur.execute("""
        SELECT body_text FROM email.messages
        WHERE to_addr::text ILIKE '%%dean@university.edu%%'
        AND subject ILIKE '%%Research%%Teaching%%Integration%%Summary%%'
        ORDER BY id DESC LIMIT 1
    """)
    row = cur.fetchone()
    if not row:
        check("Dean email body has correct faculty count, average and review count",
              False, "No dean email found")
        return
    body = (row[0] or "")
    import re as _re
    nums = [float(x) for x in _re.findall(r"\d+(?:[.,]\d+)?", body.replace(",", "."))]
    exp_count = FACULTY_COUNT
    exp_avg = sum(vals) / len(vals) if vals else 0
    exp_below = sum(1 for v in vals if v < 30)
    has_count = any(abs(n - exp_count) < 0.5 for n in nums)
    has_avg = any(abs(n - exp_avg) <= 1.0 for n in nums)
    has_below = any(abs(n - exp_below) < 0.5 for n in nums)
    check("Dean email body has correct faculty count, average and review count",
          has_count and has_avg and has_below,
          f"body nums={nums}; want count={exp_count}, avg~{exp_avg:.1f}, below={exp_below}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--agent_workspace", required=False, default=".")
    parser.add_argument("--groundtruth_workspace", required=False, default=".")
    parser.add_argument("--launch_time", required=False)
    parser.add_argument("--res_log_file", required=False)
    args = parser.parse_args()

    ws = args.agent_workspace

    conn = psycopg2.connect(**DB_CONFIG)
    cur = conn.cursor()

    # 1. Check course_keywords.json exists and has content
    kw_path = os.path.join(ws, "course_keywords.json")
    check("course_keywords.json exists", os.path.exists(kw_path), "File not found")
    if os.path.exists(kw_path):
        with open(kw_path) as f:
            kw_data = json.load(f)
        check("course_keywords has entries", len(kw_data) > 0, f"Got {len(kw_data)} entries")
    else:
        kw_data = {}

    # 2. Check alignment_scores.json exists and has correct faculty entries
    as_path = os.path.join(ws, "alignment_scores.json")
    check("alignment_scores.json exists", os.path.exists(as_path), "File not found")
    if os.path.exists(as_path):
        with open(as_path) as f:
            as_data = json.load(f)
        # Could be list or dict
        if isinstance(as_data, list):
            check(f"alignment_scores has {FACULTY_COUNT} entries", len(as_data) == FACULTY_COUNT, f"Got {len(as_data)}")
            names_in_scores = [e.get("faculty_name", e.get("name", "")) for e in as_data]
        elif isinstance(as_data, dict):
            check(f"alignment_scores has {FACULTY_COUNT} entries", len(as_data) == FACULTY_COUNT, f"Got {len(as_data)}")
            names_in_scores = list(as_data.keys())
        else:
            names_in_scores = []
            check(f"alignment_scores has {FACULTY_COUNT} entries", False, "Unexpected format")

        # Check each faculty name appears
        for fac in FACULTY:
            found = any(fac.lower() in n.lower() for n in names_in_scores)
            check(f"alignment_scores contains {fac}", found, f"Names: {names_in_scores}")
    else:
        as_data = {}

    # 3. Check summary_stats.json
    ss_data = {}
    ss_path = os.path.join(ws, "summary_stats.json")
    check("summary_stats.json exists", os.path.exists(ss_path), "File not found")
    if os.path.exists(ss_path):
        with open(ss_path) as f:
            ss_data = json.load(f)
        for key in ["total_faculty", "average_alignment_score", "max_score", "min_score"]:
            alt_keys = [key, key.replace("_", " "), key.replace("alignment_", "")]
            found = any(k in str(ss_data).lower() for k in [key.lower(), key.replace("_", "").lower()])
            # More flexible: check key exists in some form
            found = any(key.replace("_", "") in k.replace("_", "").lower() for k in ss_data.keys()) or key in ss_data
            check(f"summary_stats has {key}", found, f"Keys: {list(ss_data.keys())}")

    # 4. Check Word document exists
    doc_path = os.path.join(ws, "Research_Teaching_Report.docx")
    check("Research_Teaching_Report.docx exists", os.path.exists(doc_path), "File not found")
    if os.path.exists(doc_path):
        try:
            from docx import Document
            doc = Document(doc_path)
            full_text = "\n".join([p.text for p in doc.paragraphs])
            text_lower = full_text.lower()

            check("Report has Executive Summary",
                  any(s in text_lower for s in ["executive summary", "сводка", "резюме"]),
                  "Section not found")
            check("Report mentions alignment score",
                  ("alignment" in text_lower or "согласован" in text_lower)
                  and ("score" in text_lower or "балл" in text_lower),
                  "Not found")
            for fac in FACULTY:
                check(f"Report mentions {fac}", fac.lower() in text_lower, "Not found")
            for dept in DEPARTMENTS:
                check(f"Report mentions {dept} dept", dept.lower() in text_lower, "Not found")
            check("Report has recommendations section",
                  any(s in text_lower for s in ["recommend", "рекомендац"]),
                  "Not found")
            check("Report mentions review needed concept",
                  any(s in text_lower for s in ["review", "пересмотр", "review needed"]),
                  "Not found")
        except Exception as e:
            check("Report content readable", False, str(e))

    # 5. Check Google Sheet
    cur.execute("SELECT id FROM gsheet.spreadsheets WHERE title ILIKE '%Research_Teaching_Alignment%' OR title ILIKE '%Research%Teaching%Alignment%' LIMIT 1")
    row = cur.fetchone()
    check("GSheet spreadsheet exists", row is not None, "No matching spreadsheet found")
    if row:
        ss_id = row[0]
        cur.execute("SELECT id FROM gsheet.sheets WHERE spreadsheet_id = %s AND title ILIKE '%%Alignment%%Matrix%%' LIMIT 1", (ss_id,))
        sheet_row = cur.fetchone()
        check("GSheet has Alignment Matrix sheet", sheet_row is not None, "Sheet not found")
        if sheet_row:
            sheet_id = sheet_row[0]
            # Check header row
            cur.execute("""
                SELECT value FROM gsheet.cells
                WHERE spreadsheet_id = %s AND sheet_id = %s AND row_index = 0
                ORDER BY col_index
            """, (ss_id, sheet_id))
            headers = [r[0].lower() if r[0] else "" for r in cur.fetchall()]
            check("GSheet has Faculty Name header", any("faculty" in h and "name" in h for h in headers), f"Headers: {headers}")
            check("GSheet has Alignment Score header", any("alignment" in h or "score" in h for h in headers), f"Headers: {headers}")
            check("GSheet has Status header", any("status" in h for h in headers), f"Headers: {headers}")

            # Check data rows exist (should be 4 faculty)
            cur.execute("""
                SELECT COUNT(DISTINCT row_index) FROM gsheet.cells
                WHERE spreadsheet_id = %s AND sheet_id = %s AND row_index > 0
            """, (ss_id, sheet_id))
            data_rows = cur.fetchone()[0]
            check(f"GSheet has {FACULTY_COUNT} data rows", data_rows == FACULTY_COUNT, f"Got {data_rows} rows")

            # Check faculty names appear in the sheet
            cur.execute("""
                SELECT value FROM gsheet.cells
                WHERE spreadsheet_id = %s AND sheet_id = %s AND row_index > 0
            """, (ss_id, sheet_id))
            all_values = " ".join([r[0] for r in cur.fetchall() if r[0]]).lower()
            for fac in FACULTY:
                check(f"GSheet contains {fac}", fac.lower() in all_values, "Not found")

    # 6. Check emails sent
    # Dean email
    cur.execute("""
        SELECT id, subject, body_text FROM email.messages
        WHERE to_addr::text ILIKE '%dean@university.edu%'
        AND subject ILIKE '%Research%Teaching%Integration%Summary%'
    """)
    dean_emails = cur.fetchall()
    check("Dean email sent", len(dean_emails) >= 1, f"Found {len(dean_emails)} matching emails")
    if dean_emails:
        body = (dean_emails[0][2] or "").lower()
        check("Dean email mentions alignment", "alignment" in body or "score" in body, "Not found in body")

    # Department head emails
    cur.execute("""
        SELECT id, to_addr, subject, body_text FROM email.messages
        WHERE subject ILIKE '%Department%Research%Teaching%Alignment%Update%'
    """)
    dept_emails = cur.fetchall()
    check("Department head emails sent", len(dept_emails) >= 2, f"Found {len(dept_emails)} dept emails")

    # Check dept head recipients
    if dept_emails:
        all_recipients = " ".join([str(e[1]) for e in dept_emails]).lower()
        check("Email to biochem head", "head_biochem@university.edu" in all_recipients, f"Recipients: {all_recipients}")
        check("Email to bioinfo head", "head_bioinfo@university.edu" in all_recipients, f"Recipients: {all_recipients}")

    # ===== CRITICAL SEMANTIC CHECKS =====
    scores = check_score_semantics(as_data)
    check_status_thresholding(cur, scores)
    check_summary_semantics(ss_data, scores)
    check_email_routing_semantics(cur, scores)
    check_dean_email_semantics(cur, scores)

    cur.close()
    conn.close()

    check_reverse_validation()
    check_xlsx_content(ws)

    total = PASS_COUNT + FAIL_COUNT
    accuracy = PASS_COUNT / total * 100 if total > 0 else 0
    print(f"\nOverall: {PASS_COUNT}/{total} ({accuracy:.1f}%)")

    critical_failed = [n for n in FAILED_NAMES if n in CRITICAL_CHECKS]
    if critical_failed:
        print(f"CRITICAL FAILURES: {len(critical_failed)}")
        for n in critical_failed:
            print(f"  - {n}")

    result = {
        "total_passed": PASS_COUNT,
        "total_checks": total,
        "accuracy": accuracy,
        "critical_failed": critical_failed,
    }
    if args.res_log_file:
        with open(args.res_log_file, "w") as f:
            json.dump(result, f, indent=2)

    if critical_failed:
        print("FAIL (critical check failed)")
        sys.exit(1)
    sys.exit(0 if accuracy >= 70 else 1)


if __name__ == "__main__":
    main()
