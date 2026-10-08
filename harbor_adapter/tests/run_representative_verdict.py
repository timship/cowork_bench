"""Run four real graders through pass, content fail, and a crash.

Pass for the DB-backed graders uses a local fixture for the database cursor.
The grader's own checks and its own final branch still produce the Pass line.
Stdout is not rewritten.
"""

from __future__ import annotations

import importlib.util
import json
import runpy
import shutil
import sys
import tempfile
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
TASKS = ROOT / "tasks" / "finalpool"


def _load(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def _counts(text: str) -> dict[str, int]:
    true = false = 0
    for line in text.splitlines():
        if line.strip() == "Pass: True":
            true += 1
        elif line.strip() == "Pass: False":
            false += 1
    return {"pass_true": true, "pass_false": false}


def _run_main(module, argv: list[str]) -> tuple[int, str]:
    buffer: list[str] = []
    original = sys.stdout

    class _Capture:
        def write(self, data: str) -> int:
            buffer.append(data)
            return len(data)

        def flush(self) -> None:
            return None

    capture = _Capture()
    sys.argv = argv
    sys.stdout = capture
    code = 0
    try:
        module.main()
    except SystemExit as exc:
        code = exc.code if isinstance(exc.code, int) else 1
    except Exception:
        traceback.print_exc(file=capture)
        code = 1
    finally:
        sys.stdout = original
    return code, "".join(buffer)


class _Cursor:
    def __init__(self, handler):
        self.handler = handler
        self.sql = ""
        self.params = ()

    def execute(self, sql, params=None):
        self.sql = sql
        self.params = params or ()

    def fetchall(self):
        return self.handler(self.sql, self.params, "all")

    def fetchone(self):
        return self.handler(self.sql, self.params, "one")

    def close(self):
        return None


class _Conn:
    def __init__(self, handler):
        self.handler = handler

    def cursor(self):
        return _Cursor(self.handler)

    def close(self):
        return None


def _patch_connect(module, handler):
    original = module.psycopg2.connect

    def connect(**_kwargs):
        return _Conn(handler)

    module.psycopg2.connect = connect
    return original


def _restore_connect(module, original) -> None:
    module.psycopg2.connect = original


def _run_path(path: Path, argv: list[str]) -> tuple[int, str]:
    buffer: list[str] = []
    original = sys.stdout

    class _Capture:
        def write(self, data: str) -> int:
            buffer.append(data)
            return len(data)

        def flush(self) -> None:
            return None

    capture = _Capture()
    sys.argv = argv
    sys.stdout = capture
    code = 0
    try:
        runpy.run_path(str(path), run_name="__main__")
    except SystemExit as exc:
        code = exc.code if isinstance(exc.code, int) else 1
    except Exception:
        traceback.print_exc(file=capture)
        code = 1
    finally:
        sys.stdout = original
    return code, "".join(buffer)


def _xlsx(path: Path, sheets: dict[str, list[list]]) -> None:
    import openpyxl

    book = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = book.active if first else book.create_sheet(name)
        sheet.title = name
        first = False
        for row in rows:
            sheet.append(row)
    book.save(path)


def _pptx(path: Path, slides: list[str]) -> None:
    from pptx import Presentation

    presentation = Presentation()
    for text in slides:
        layout = presentation.slide_layouts[1] if len(presentation.slide_layouts) > 1 else presentation.slide_layouts[0]
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = text
    presentation.save(path)


def _docx(path: Path, paragraphs: list[str]) -> None:
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(path)


def insales_pass(workspace: Path) -> None:
    categories = ["ТВ и домашний кинотеатр", "Электроника", "Аудио", "Камеры", "Часы", "Бытовая техника"]
    performance = [
        ["Category", "Revenue", "Units_Sold", "Avg_Rating"],
        *[[name, 100, 10, 4.54] for name in categories],
    ]
    market = [["Category", "Market_Growth_Rate"]]
    opportunities = [["Category", "Growth", "Rating", "Opportunity_Score"]]
    market_data = {
        "ТВ и домашний кинотеатр": 12.5,
        "Электроника": 8.3,
        "Аудио": 15.2,
        "Камеры": 3.1,
        "Часы": 6.7,
        "Бытовая техника": -2.4,
    }
    for name, growth in market_data.items():
        market.append([name, growth])
        score = round(growth * 4.54, 2)
        opportunities.append([name, growth, 4.54, score])
    opportunities[1], opportunities[3] = opportunities[3], opportunities[1]
    _xlsx(
        workspace / "Product_Review.xlsx",
        {
            "Category Performance": performance,
            "Market Comparison": market,
            "Opportunities": opportunities,
        },
    )
    _pptx(
        workspace / "Q1_Product_Review.pptx",
        ["product performance revenue", "audio market trend", "opportunity one", "opportunity two", "opportunity three"],
    )


def insales_handler(sql, _params, kind):
    if kind == "all" and "revenue" in sql.lower():
        return [("Аудио", 100, 10, 10)]
    if kind == "all":
        return [("Аудио", 4.54)]
    return None


def canvas_pass(workspace: Path) -> None:
    _xlsx(
        workspace / "Faculty_Workload.xlsx",
        {
            "Instructor Load": [
                ["instructor", "courses_count", "total_students", "overloaded_yn"],
                ["Д-р Андрей Волков", 3, 2728, "yes"],
                ["Д-р Аделина Мартынова", 1, 365, "no"],
            ],
            "Department Summary": [["department"]] + [[f"dept {i}"] for i in range(7)],
        },
    )
    _pptx(workspace / "Workload_Review.pptx", ["workload review", "details", "summary of overload"])


def canvas_handler(sql, _params, kind):
    text = sql.lower()
    if "teacherenrollment" in text:
        return [
            ("Д-р Андрей Волков", 3, 2728, 1, 1),
            ("Д-р Аделина Мартынова", 1, 365, 1, 1),
        ]
    if "gsheet.spreadsheets" in text:
        return [(1, "Faculty Workload")]
    if "count(*)" in text:
        return (11,)
    return [] if kind == "all" else None


def terminal_cells():
    recipes = [
        "Борщ", "Оливье", "Блины", "Пельмени", "Сырники",
        "Уха", "Солянка", "Винегрет", "Голубцы",
    ]
    recipe = {}
    headers = ["recipe_name", "meal_type", "calories", "protein", "carbs", "fat", "fiber", "sodium"]
    for col, header in enumerate(headers):
        recipe[(0, col)] = header
    for row, name in enumerate(recipes, start=1):
        recipe[(row, 0)] = name
        for col in range(1, len(headers)):
            recipe[(row, col)] = 10
    daily = {
        (0, 0): "meal",
        (0, 1): "calories",
        (0, 2): "pct_daily_calories",
        (1, 0): "breakfast",
        (1, 1): 500,
        (2, 0): "lunch",
        (2, 1): 700,
        (3, 0): "dinner",
        (3, 1): 800,
        (4, 0): "total",
        (4, 1): 2000,
        (4, 2): 100,
        (5, 0): "standards 50 300 65 25 2300",
    }
    return {10: recipe, 11: daily}


def terminal_handler(sql, params, kind):
    text = " ".join(sql.lower().split())
    cells = terminal_cells()
    if "from gsheet.spreadsheets" in text and "group by" in text:
        return []
    if "from gsheet.spreadsheets" in text:
        return [(1, "Nutrition Dashboard")]
    if "from gsheet.sheets" in text:
        return [(10, "Recipe Comparison"), (11, "Daily Plan")]
    if "from gsheet.cells" in text and "row_index = 0" in text:
        sheet_cells = cells[10]
        return [(sheet_cells[(0, col)],) for col in range(8)]
    if "from gsheet.cells" in text:
        sheet_id = params[1] if len(params) > 1 else 10
        grid = cells.get(sheet_id, {})
        return [(row, col, value) for (row, col), value in sorted(grid.items())]
    return [] if kind == "all" else None


def terminal_pass(workspace: Path) -> None:
    _docx(
        workspace / "Wellness_Diet_Plan.docx",
        [
            "Wellness diet plan and recommended standards",
            "Breakfast lunch dinner meal plan",
            "Daily plan deficit supplement",
            "Calories 2000 protein 50 carbs 300 fat 65 fiber 25 sodium 2300",
        ],
    )
    (workspace / "nutrition_calculator.py").write_text("print(1)\n", encoding="utf-8")


def kulinar_rows():
    overview = [
        ["recipe_name", "category", "difficulty", "estimated_calories", "protein_level", "fiber_level"],
        ["Борщ", "soup", "easy", 77, "high", "high"],
        ["Оливье", "salad", "medium", 250, "high", "low"],
        ["Блины", "bakery", "easy", 165, "medium", "low"],
        ["Пельмени", "main", "hard", 130, "high", "low"],
        ["Сырники", "dessert", "easy", 60, "medium", "low"],
    ]
    scored = []
    points = {"high": 3, "medium": 2, "low": 1}
    for row in overview[1:]:
        score = (points[row[4]] + points[row[5]]) * 10 - row[3] / 100.0
        scored.append((score, row[0]))
    scored.sort(reverse=True)
    recommendation = [["rank", "recipe_name", "health_score", "reason"]]
    for rank, (score, name) in enumerate(scored, start=1):
        recommendation.append([rank, name, round(score, 2), "highest remaining health score"])
    return overview, recommendation, [name for _, name in scored]


def kulinar_pass(workspace: Path, guide: Path) -> None:
    overview, recommendation, names = kulinar_rows()
    _xlsx(
        workspace / "Nutrition_Comparison.xlsx",
        {"Recipe Overview": overview, "Recommendation": recommendation},
    )
    slides = ["Healthy Eating Guide"] + [f"Recipe {name}" for name in names[:3]] + ["Summary of the eating guide"]
    while len(slides) < 7:
        slides.append("Additional nutrition context for the guide")
    _pptx(workspace / "Healthy_Eating_Guide.pptx", slides)
    shutil.copy(guide, workspace / "nutrition_guide.md")


def main() -> int:
    out = Path(sys.argv[1])
    out.mkdir(parents=True, exist_ok=True)
    results = []
    with tempfile.TemporaryDirectory(prefix="verdict-repr-") as tmp:
        base = Path(tmp)

        def record(task: str, mode: str, code: int, text: str) -> None:
            counts = _counts(text)
            safe = [
                line for line in text.splitlines()
                if "password" not in line.lower() and "pgpassword" not in line.lower()
            ]
            results.append(
                {
                    "task": task,
                    "mode": mode,
                    "exit": code,
                    **counts,
                    "traceback": "Traceback (most recent call last):" in text,
                    "tail": safe[-12:],
                }
            )
            print(f"{task} {mode} exit={code} true={counts['pass_true']} false={counts['pass_false']} traceback={'yes' if results[-1]['traceback'] else 'no'}")

        # insales
        task = "insales-product-launch-dashboard"
        path = TASKS / task / "evaluation" / "main.py"
        empty = base / "insales-empty"
        empty.mkdir()
        module = _load(path, "insales_fail")
        code, text = _run_main(module, ["main.py", "--agent_workspace", str(empty)])
        record(task, "fail", code, text)
        good = base / "insales-pass"
        good.mkdir()
        insales_pass(good)
        module = _load(path, "insales_pass")
        original = _patch_connect(module, insales_handler)
        try:
            code, text = _run_main(module, ["main.py", "--agent_workspace", str(good)])
        finally:
            _restore_connect(module, original)
        record(task, "pass", code, text)
        crash_ws = base / "insales-crash"
        crash_ws.mkdir()
        _xlsx(crash_ws / "Product_Review.xlsx", {"Sheet": [["a"]]})
        module = _load(path, "insales_crash")
        code, text = _run_main(module, ["main.py", "--agent_workspace", str(crash_ws)])
        record(task, "crash", code, text)

        # canvas
        task = "canvas-faculty-workload-review"
        path = TASKS / task / "evaluation" / "main.py"
        empty = base / "canvas-empty"
        empty.mkdir()
        module = _load(path, "canvas_fail")
        code, text = _run_main(module, ["main.py", "--agent_workspace", str(empty)])
        record(task, "fail", code, text)
        good = base / "canvas-pass"
        good.mkdir()
        canvas_pass(good)
        module = _load(path, "canvas_pass")
        original = _patch_connect(module, canvas_handler)
        try:
            code, text = _run_main(module, ["main.py", "--agent_workspace", str(good)])
        finally:
            _restore_connect(module, original)
        record(task, "pass", code, text)
        crash_ws = base / "canvas-crash"
        crash_ws.mkdir()
        _xlsx(crash_ws / "Faculty_Workload.xlsx", {"Instructor Load": [["instructor"]]})
        module = _load(path, "canvas_crash")
        code, text = _run_main(module, ["main.py", "--agent_workspace", str(crash_ws)])
        record(task, "crash", code, text)

        # terminal
        task = "terminal-kulinar-pw-nutrition-gsheet-word"
        path = TASKS / task / "evaluation" / "main.py"
        empty = base / "terminal-empty"
        empty.mkdir()
        module = _load(path, "terminal_fail")

        def empty_db(sql, params, kind):
            return [] if kind == "all" else None

        original = _patch_connect(module, empty_db)
        try:
            code, text = _run_main(module, ["main.py", "--agent_workspace", str(empty)])
        finally:
            _restore_connect(module, original)
        record(task, "fail", code, text)
        good = base / "terminal-pass"
        good.mkdir()
        terminal_pass(good)
        module = _load(path, "terminal_pass")
        original = _patch_connect(module, terminal_handler)
        try:
            code, text = _run_main(module, ["main.py", "--agent_workspace", str(good)])
        finally:
            _restore_connect(module, original)
        record(task, "pass", code, text)
        module = _load(path, "terminal_crash")
        code, text = _run_main(module, ["main.py", "--agent_workspace", str(empty)])
        record(task, "crash", code, text)

        # kulinar
        task = "kulinar-nutrition-ppt"
        path = TASKS / task / "evaluation" / "main.py"
        guide = TASKS / task / "groundtruth_workspace" / "nutrition_guide.md"
        empty = base / "kulinar-empty"
        empty.mkdir()
        code, text = _run_path(path, ["main.py", "--agent_workspace", str(empty)])
        record(task, "fail", code, text)
        good = base / "kulinar-pass"
        good.mkdir()
        kulinar_pass(good, guide)
        code, text = _run_path(path, ["main.py", "--agent_workspace", str(good)])
        record(task, "pass", code, text)
        code, text = _run_path(
            path,
            ["main.py", "--agent_workspace", str(good), "--res_log_file", str(good)],
        )
        record(task, "crash", code, text)

    (out / "representative_results.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    failed = [
        item for item in results
        if (item["mode"] == "pass" and not (item["exit"] == 0 and item["pass_true"] == 1 and item["pass_false"] == 0))
        or (item["mode"] == "fail" and not (item["exit"] == 1 and item["pass_false"] == 1 and item["pass_true"] == 0 and not item["traceback"]))
        or (item["mode"] == "crash" and not (item["pass_true"] == 0 and item["pass_false"] == 0 and item["traceback"]))
    ]
    print("FAILED", failed)
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
